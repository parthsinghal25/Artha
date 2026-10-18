"""IP01 (owned-folder binding) and IP02 (D29 effective asset ruling) from the case-11 integration-prep review.

Every ledger here is a disposable fixture: pack set 5 is built into a temporary folder, and pack set 4 is a temporary COPY. The real
ledgers, packs, accepted packages and review records are only read. Card returns are structured fixture text, not source prose."""
import copy
import io
import json
import os
import shutil
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import pytest

from conftest import write_decision
from sdgen import contract, effective_spec, progress, workpack4, workpack5
from sdgen.authoring_check import author_check, load_pack
from sdgen.card_adapter import native_cells, render_card
from sdgen.contract import approved_clarifications, derive_requirements
from sdgen.microbatch import build_package, verify_package
from sdgen.paths import ROOT
from sdgen.progress import GateRefused, LedgerIntegrityError

STAGE = "remaining_cards"
D29_ASSET = "AST-eeb19c880495b2ead3"
EXCLUDED = ["WORK-R2-85e527fa30fb92db46", "WORK-R2-f8ded04181fbc6806b"]
RULINGS = json.loads(workpack5.RULINGS.read_text())
EFFORT = next(c for c in RULINGS["clarifications"] if c["id"] == "CL11-EFFORT-02")


def later(days=1):
    t = datetime.now(timezone.utc) + timedelta(days=days)
    return t, t.strftime("%Y-%m-%dT%H:%M:%SZ")


def style(card, profile="S0", features=()):
    return {"card": card, "applied_profile": profile, "features_used": list(features), "typos_used": [], "protected_values_checked": True, "concerns": []}


def ledger_return(card):
    return {"card": card, "revision": 1, "kind": "ledger_export", "style_completion": style(card)}


def fixture_return(pack, card):
    """A structured fixture return for a prose card: one paragraph made of the card's own derived fixed values only."""
    plan, sty, prof, _, clars, _ = load_pack(pack)
    art = plan.artifact(card)
    kind = contract.allowed_kinds(art)[0][0]
    if kind == "ledger_export":
        return ledger_return(card)
    named = [plan.assets[a]["filename"] for a in art.get("attachments", []) if plan.assets[a].get("delivery_state") == "held"]
    req = derive_requirements(plan, art, kind, approved_clarifications(clars), named if kind in contract.EMAIL_KINDS else None)
    toks = []
    for g in req.required:
        toks.append(f"A${g.minor / 100:,.2f}" if g.minor is not None else g.label.split()[-1])
    row = next(r for r in sty["cards"] if r["artifact_ref"] == card)
    feats = () if prof[row["profile_id"]].get("difficulty") == "clean" else ("fixture feature",)
    d = {"card": card, "revision": 1, "kind": kind, "blocks": [{"p": "Fixture text: " + "; ".join(toks) + "."}],
         "style_completion": style(card, row["profile_id"], feats)}
    if kind in contract.EMAIL_KINDS:
        d.update(greeting="Hello,", signoff=["Fixture"])
    return d


def put(folder, d):
    folder.mkdir(parents=True, exist_ok=True)
    f = folder / f"{d['card']}.json"
    f.write_text(json.dumps(d, indent=1))
    return f


@pytest.fixture(scope="module")
def built5(tmp_path_factory):
    root = tmp_path_factory.mktemp("ps5") / "worker-packs-v5"
    entry = workpack5.build(root)
    return root, entry


@pytest.fixture()
def ps5(built5, tmp_path):
    """A private copy of pack set 5 with its outputs folder, and a fresh ledger initialised from its index."""
    root, entry = built5
    dst = tmp_path / "synthesis/worker-packs-v5"
    shutil.copytree(root, dst)
    shutil.copytree(root.parent / "worker-outputs-v5", tmp_path / "synthesis/worker-outputs-v5")
    progress.init(dst)
    return dst


@pytest.fixture()
def pack5(built5):
    root, entry = built5
    return root / entry["pack"]


@pytest.fixture()
def reviews(tmp_path):
    d = tmp_path / "reviews"
    d.mkdir()
    return d


def stage_cards(root):
    pack = next(root.glob("11-C-*"))
    return next(s for s in json.loads((pack / "STAGE.json").read_text())["stages"] if s["stage"] == STAGE)["cards"]


def returns_into(folder, cards):
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "11-D19.json").write_text(json.dumps(ledger_return("11-D19")))
    (folder / "RETURN-CHECKLIST.json").write_text(json.dumps({"case": 11, "cards": [{"card": c, "status": "fixture", "questions": []} for c in cards]}))


# ================================================================== IP01
def test_ip01_v5_ledger_binds_the_folder_its_index_declares(ps5):
    state = json.loads((ps5 / "PROGRESS.json").read_text())
    assert state["outputs_binding"]["outputs_dirname"] == "worker-outputs-v5"
    assert state["outputs_binding"]["packs_index_sha256"]
    want = ps5.parent / "worker-outputs-v5" / state["jobs"][0]["pack"] / STAGE
    assert progress.owned_folder_for(11, STAGE, ps5) == want


def test_ip01_real_lifecycle_hashes_the_bound_folder_and_refuses_changed_output(ps5, reviews):
    job = json.loads((ps5 / "PROGRESS.json").read_text())["jobs"][0]
    v5 = ps5.parent / "worker-outputs-v5" / job["pack"] / STAGE
    decoy = ps5.parent / "worker-outputs-v3" / job["pack"] / STAGE  # what the defective resolver used to read
    returns_into(decoy, stage_cards(ps5))
    progress.authorise(11, STAGE, write_decision(reviews, ps5, 11, STAGE), ps5, reviews_root=reviews)
    st = progress.set_state(11, STAGE, "working", ps5, worker="fixture", reviews_root=reviews)
    with pytest.raises(GateRefused, match="worker-outputs-v5"):
        progress.set_state(11, STAGE, "submitted", ps5)  # the v3-named decoy holds files, but the bound v5 folder holds none
    returns_into(v5, stage_cards(ps5))
    st = progress.set_state(11, STAGE, "submitted", ps5)
    assert st["output_digest"] == progress.tree_digest(v5) != progress.tree_digest(decoy)
    events = [json.loads(l) for l in (ps5 / "PROGRESS-EVENTS.jsonl").read_text().splitlines()]
    assert [e["detail"].get("owned_folder") for e in events if e["action"] == "set_state"] == [str(v5), str(v5)]
    # changed output after submission: acceptance refused and routed to repair, even with a record bound to the submitted digest
    (v5 / "11-D19.json").write_text(json.dumps({**ledger_return("11-D19"), "revision": 2}))
    now, at = later()
    rec = write_decision(reviews, ps5, 11, STAGE, kind="authoring_stage_acceptance", decision="accepted", output_digest=st["output_digest"],
                         name="acc1.json", decided_at=at)
    with pytest.raises(GateRefused, match="differs from the submitted digest"):
        progress.set_state(11, STAGE, "independently_accepted", ps5, acceptance_record=rec, reviews_root=reviews, now=now + timedelta(hours=1))
    assert json.loads((ps5 / "PROGRESS.json").read_text())["jobs"][0]["stages"][STAGE]["state"] == "repair_required"
    # a change in the decoy folder is irrelevant; resubmission of the bound folder and a fresh record are accepted
    progress.set_state(11, STAGE, "working", ps5, reviews_root=reviews)
    st2 = progress.set_state(11, STAGE, "submitted", ps5)
    (decoy / "11-D19.json").write_text("{}")
    rec2 = write_decision(reviews, ps5, 11, STAGE, kind="authoring_stage_acceptance", decision="accepted", output_digest=st2["output_digest"],
                          name="acc2.json", decided_at=at)
    res = progress.set_state(11, STAGE, "independently_accepted", ps5, acceptance_record=rec2, reviews_root=reviews, now=now + timedelta(hours=1))
    assert res["acceptance"]["output_digest"] == progress.tree_digest(v5) and progress.verify_ledger(ps5)


def test_ip01_legacy_v3_ledger_still_resolves_its_original_folder(pack_set_copy):
    progress.init(pack_set_copy)
    state = json.loads((pack_set_copy / "PROGRESS.json").read_text())
    assert state["outputs_binding"] == {"outputs_dirname": "worker-outputs-v3", "source": "legacy pack-set table (index predates the binding)",
                                        "packs_index_sha256": state["outputs_binding"]["packs_index_sha256"]}
    job = next(j for j in state["jobs"] if j["case"] == 11)
    assert progress.owned_folder(pack_set_copy, job, "style_microbatch") == pack_set_copy.parent / "worker-outputs-v3" / job["pack"] / "style_microbatch"


@pytest.mark.parametrize("version", [3, 4])
def test_ip01_existing_ledgers_without_a_binding_resolve_by_their_pack_set(tmp_path, version):
    """Copies of the real pack-3 and pack-4 ledgers (created before the binding existed) resolve their own folders; the originals are only read."""
    src = ROOT / f"Data/synthesis/worker-packs-v{version}"
    dst = tmp_path / f"worker-packs-v{version}"
    dst.mkdir()
    for name in ("PROGRESS.json", "PROGRESS-EVENTS.jsonl", "PACKS-INDEX.json"):
        shutil.copy(src / name, dst / name)
    state = json.loads((dst / "PROGRESS.json").read_text())
    assert "outputs_binding" not in state and progress.verify_ledger(dst)
    job = next(j for j in state["jobs"] if j["case"] == 11)
    stage = "style_microbatch" if version == 3 else STAGE
    assert progress.owned_folder(dst, job, stage) == tmp_path / f"worker-outputs-v{version}" / job["pack"] / stage


def test_ip01_index_that_contradicts_the_ledger_binding_is_refused(ps5, reviews):
    progress.authorise(11, STAGE, write_decision(reviews, ps5, 11, STAGE), ps5, reviews_root=reviews)
    idx = json.loads((ps5 / "PACKS-INDEX.json").read_text())
    (ps5 / "PACKS-INDEX.json").write_text(json.dumps({**idx, "outputs_dirname": "worker-outputs-v3"}))
    with pytest.raises(GateRefused, match="binding mismatch"):
        progress.set_state(11, STAGE, "working", ps5, reviews_root=reviews)


def test_ip01_editing_the_binding_in_the_ledger_breaks_its_integrity(ps5):
    state = json.loads((ps5 / "PROGRESS.json").read_text())
    state["outputs_binding"]["outputs_dirname"] = "worker-outputs-v3"
    (ps5 / "PROGRESS.json").write_text(json.dumps(state))
    with pytest.raises(LedgerIntegrityError):
        progress.owned_folder_for(11, STAGE, ps5)


@pytest.mark.parametrize("bad", ["../escape", "worker-outputs-v5/../../x", "/tmp/worker-outputs", "outputs"])
def test_ip01_an_unsafe_declared_folder_is_refused_at_init(built5, tmp_path, bad):
    root, _ = built5
    dst = tmp_path / "worker-packs-v5"
    shutil.copytree(root, dst)
    idx = json.loads((dst / "PACKS-INDEX.json").read_text())
    (dst / "PACKS-INDEX.json").write_text(json.dumps({**idx, "outputs_dirname": bad}))
    with pytest.raises(GateRefused, match="no valid owned-output folder binding"):
        progress.init(dst)
    assert not (dst / "PROGRESS.json").exists()


def test_ip01_an_owned_folder_that_escapes_through_a_symlink_is_refused(ps5, reviews, tmp_path):
    job = json.loads((ps5 / "PROGRESS.json").read_text())["jobs"][0]
    owned = ps5.parent / "worker-outputs-v5" / job["pack"] / STAGE
    outside = tmp_path / "elsewhere"
    returns_into(outside, stage_cards(ps5))
    shutil.rmtree(owned)
    os.symlink(outside, owned)
    progress.authorise(11, STAGE, write_decision(reviews, ps5, 11, STAGE), ps5, reviews_root=reviews)
    with pytest.raises(GateRefused, match="resolves outside its bound outputs folder"):
        progress.set_state(11, STAGE, "working", ps5, reviews_root=reviews)


# ================================================================== IP02
def d29_rendered(pack):
    plan = load_pack(pack)[0]
    rc = render_card(ledger_return("11-D29"), plan, plan.artifact("11-D29"))
    [(rel, data)] = rc.files.items()
    return plan, rel, data


def test_ip02_d29_renders_a_real_xlsx_with_exactly_the_ruled_rows(pack5):
    plan, rel, data = d29_rendered(pack5)
    assert rel == "originals/ledgers/11-16-complaints-effort-export.xlsx" and data[:2] == b"PK"
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    [ws] = wb.worksheets
    cells = [list(r) for r in ws.iter_rows(values_only=True)]
    header, body = cells[0], cells[1:]
    assert header == ["export_date", "case_id", "work_id", "event_date", "actor", "seconds", "phase"]
    assert len(body) == 22 and sum(r[5] for r in body) == 19920 and all(type(r[5]) is int for r in body)
    expected = EFFORT["effective_asset_projection"]["expected_rows"]
    assert [dict(zip(header[2:], r[2:])) for r in body] == expected
    assert all(r[:2] == ["2026-03-27", "C-e1da28538652e263887a"] for r in body)
    flat = native_cells(data, plan.assets[D29_ASSET]["media_type"])
    assert not any(contract.STAGING_CARD_REF.search(c) for c in flat) and not any(w in flat for w in EXCLUDED)
    assert not any("EXT-AFCA" in c or "related_artifact" in c for c in flat)
    assert {r[6] for r in body} == {"complaint_handling", "external_dispute"}  # broad raw phases kept
    assert wb.properties.created == datetime(2026, 3, 27, 23, 0)  # the export's own creation time, not a historical work date


def test_ip02_retained_cells_equal_the_frozen_plan_cells(pack5):
    frozen = json.loads((pack5 / "plan/assistant-pack.json").read_text())
    planned = next(a for a in frozen["assets"] if a["asset_id"] == D29_ASSET)
    kept = [{k: v for k, v in r.items() if k != "related_artifact"} for r in planned["document_spec"]["rows"] if r["work_id"] not in EXCLUDED]
    _, _, data = d29_rendered(pack5)
    ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    assert [dict(zip(rows[0][2:], r[2:])) for r in rows[1:]] == kept
    assert len(planned["document_spec"]["rows"]) == 24  # the frozen plan and its 24-row history are unchanged


def test_ip02_effective_asset_keeps_identity_dates_and_reindexes_temporal_roles(pack5):
    plan = load_pack(pack5)[0]
    a = plan.assets[D29_ASSET]
    frozen = next(x for x in json.loads((pack5 / "plan/assistant-pack.json").read_text())["assets"] if x["asset_id"] == D29_ASSET)
    for k in ("asset_id", "source_id", "filename", "format", "media_type", "created_at", "available_at", "owner"):
        assert a[k] == frozen[k]
    roles = a["document_spec"]["field_temporal_roles"]
    assert len(roles) == 2 + 22 * 5 and not any("related_artifact" in p for p in roles) and "/rows/21/phase" in roles and "/rows/22/phase" not in roles
    assert plan.assets_by_name["11-16-complaints-effort-export.xlsx"] is a and len(plan.pack["work_entries"]) == 22
    eff = json.loads((pack5 / "plan/EFFECTIVE-ASSET-SPECS.json").read_text())
    assert [c["card"] for c in eff["cards_whose_recorded_internal_effort_differs_from_planned_duration"]] == ["11-D21", "11-D25"]


def test_ip02_author_check_requires_every_ruled_value_and_none_of_the_rejected(pack5, tmp_path):
    rows = author_check(put(tmp_path, ledger_return("11-D29")), pack5, STAGE)
    assert all(r["ok"] for r in rows), [r for r in rows if not r["ok"]]
    native = [r["detail"] for r in rows if r["id"] == "A6n"]
    assert len(native) == 2 + 22 * 5
    assert not any("related_artifact" in d or "EXT-AFCA" in d for d in native)


@pytest.mark.parametrize("mutate,match", [
    (lambda c: c["effective_asset_projection"]["expected_rows"][0].update(seconds=301), "expected rows"),
    (lambda c: c["effective_asset_projection"].update(expected_seconds=21720), "row count/seconds"),
    (lambda c: c["effective_asset_projection"].update(available_at="2026-01-01T00:00:00+11:00"), "availability is never moved"),
    (lambda c: c["effective_asset_projection"].update(format="CSV"), "never converts"),
    (lambda c: c["effective_asset_projection"].update(excluded_work_ids=EXCLUDED[:1]), "excluded rows differs|expected rows"),
    (lambda c: c["private_original_rows_excluded"][0].update(seconds=60), "private record of excluded rows"),
])
def test_ip02_a_ruling_that_disagrees_with_the_frozen_plan_is_refused(pack5, mutate, match):
    frozen = json.loads((pack5 / "plan/assistant-pack.json").read_text())
    c = copy.deepcopy(EFFORT)
    mutate(c)
    with pytest.raises(effective_spec.RulingMismatch, match=match):
        effective_spec.effective_assets(frozen, [c])


def test_ip02_the_staging_label_guard_is_not_waived(pack5):
    """A projection that kept the authoring column would still be refused by the renderer's unchanged guard."""
    frozen = json.loads((pack5 / "plan/assistant-pack.json").read_text())
    planned = next(a for a in frozen["assets"] if a["asset_id"] == D29_ASSET)
    plan = load_pack(pack5)[0]
    plan.assets[D29_ASSET] = planned  # the unruled asset
    with pytest.raises(Exception, match="staging card references"):
        render_card(ledger_return("11-D29"), plan, plan.artifact("11-D29"))


def test_ip02_a_tampered_or_missing_effective_spec_file_refuses_the_pack(pack5, tmp_path):
    bad = tmp_path / "bad"
    shutil.copytree(pack5, bad)
    f = bad / "plan/EFFECTIVE-ASSET-SPECS.json"
    doc = json.loads(f.read_text())
    doc["assets"][0]["rows"][0]["seconds"] = 9999
    f.write_text(json.dumps(doc))
    with pytest.raises(effective_spec.RulingMismatch, match="differs from the projection"):
        load_pack(bad)
    f.unlink()
    with pytest.raises(effective_spec.RulingMismatch, match="no plan/EFFECTIVE-ASSET-SPECS.json"):
        load_pack(bad)


def test_ip02_pack_set_5_is_deterministic_and_preserves_the_frozen_plan(built5, tmp_path):
    root, entry = built5
    again = workpack5.build(tmp_path / "worker-packs-v5")
    assert again["pack_digest"] == entry["pack_digest"] and all(c["ok"] for c in entry["checks"])
    assert entry["input_hashes"]["plan_sha256"] == RULINGS["bound_plan_sha256"] == "df33c8afd103b5b476005e9ee47deddae9af72477bb156450cd8f57946c7b6de"
    idx = json.loads((root / "PACKS-INDEX.json").read_text())
    assert idx["outputs_dirname"] == "worker-outputs-v5" and idx["predecessor"]["pack_digest"] == workpack5.PREDECESSOR["pack_digest"]
    st = json.loads((root / entry["pack"] / "STAGE.json").read_text())
    rest = next(s for s in st["stages"] if s["stage"] == STAGE)
    assert "pending_rulings" not in rest and len(rest["cards"]) == 17 and st["pack_set_version"] == 5
    with pytest.raises(FileExistsError):
        workpack5.build(root)


def test_ip02_rulings_that_differ_from_the_reviewed_hash_are_refused(monkeypatch, tmp_path):
    monkeypatch.setattr(workpack5, "RULINGS_SHA256", "0" * 64)
    with pytest.raises(RuntimeError, match="differs from the hash recorded"):
        workpack5.build(tmp_path / "x")


def test_ip02_context_stays_earlier_only_and_bound_to_the_accepted_export(pack5):
    plan = load_pack(pack5)[0]
    ctx = json.loads((pack5 / "context/ACCEPTED-PRIOR-ORIGINALS.json").read_text())
    assert ctx["bound_to_manifest"]["delivery_manifest_sha256"] == workpack4.EXPORT_MANIFEST_SHA256
    assert all(plan.artifact(c)["event_at"] > o["event_at"] for o in ctx["originals"] for c in o["usable_before_cards"])
    guide = (pack5 / "shared/WORKER-GUIDE.md").read_text()
    assert "19,920 seconds" in guide and "Pending ruling" not in guide


# ================================================================== packaging through the real adapters
def test_supported_paths_package_through_their_real_adapters(pack5, tmp_path):
    """email (11-D04), external notice (11-D21), note (11-D14) and ledger exports (11-D19, 11-D28, 11-D29) in one fixture package."""
    outs = tmp_path / "owned"
    for card in ("11-D04", "11-D21", "11-D14", "11-D19", "11-D28", "11-D29"):
        f = put(outs, fixture_return(pack5, card))
        bad = [r for r in author_check(f, pack5, STAGE) if not r["ok"]]
        assert not bad, (card, bad)
    pkg = build_package(pack5, STAGE, outs, tmp_path / "rel", "case-11-remaining-fixture", accepted_package=workpack4.EXPORT)
    man = json.loads((pkg / "intake-bundle/delivery-manifest.json").read_text())
    kinds = {o["staging_card_ref"]: (o["kind"], o["file"].split("/")[1]) for o in man["originals"] if o.get("staging_card_ref")}
    assert kinds["11-D04"] == ("email", "emails") and kinds["11-D21"] == ("email", "emails") and kinds["11-D14"] == ("note", "notes")
    led = {o["file"]: o["original_id"] for o in man["originals"] if o["file"].startswith("originals/ledgers/")}
    assert led["originals/ledgers/11-16-complaints-effort-export.xlsx"] == D29_ASSET and len(led) == 3
    assert any(o["filename"] == "11-12-external-dispute-notice.pdf" for o in man["originals"] if o["kind"] == "attachment_original")
    assert all(c["ok"] for c in verify_package(pkg, accepted_package=workpack4.EXPORT)) and man["pack_set_version"] == 5
    _, _, data = d29_rendered(pack5)
    assert (pkg / "intake-bundle/originals/ledgers/11-16-complaints-effort-export.xlsx").read_bytes() == data
