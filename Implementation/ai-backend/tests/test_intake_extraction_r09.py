"""R09: delivery/extraction/tool seam. All inputs are generated specimens in temporary directories."""
import email.message
import io
import json

import pytest

from specimens import case_row, decide, ingest_specimen, mini_bundle

CASE = "C-e1da28538652e263887a"


def _ingest(rt, bundle):
    from sdbackend.ingest.build import ingest
    from sdbackend.intake.delivery import receive
    return ingest(rt, receive(bundle, rt=rt))


def _png_with_text(text: str) -> bytes:
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (900, 200), "white")
    ImageDraw.Draw(im).text((30, 70), text, fill="black", font_size=48)
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def test_ocr_runtime_diagnostic_is_explicit(monkeypatch):
    from sdbackend.extract import image
    def unavailable(_):
        raise image.OCRUnavailable("service unreachable (test)")
    monkeypatch.setattr(image, "_ocr_lines", unavailable)
    out = image.extract_image(_png_with_text("x"))
    assert out.state == "failed" and "OCR unavailable in this runtime" in out.omissions[0]
    assert "unreachable" in out.omissions[0]


def test_receipt_image_ocr_on_supported_runtime(rehearsal_bundle):
    from sdbackend.extract.image import extract_image, ocr_runtime_status
    status = ocr_runtime_status()
    if not status["available"]:
        pytest.skip(f"OCR unavailable in this runtime (honest skip, not a pass): {status['reason']}")
    out = extract_image((rehearsal_bundle / "originals/attachments/11-05-portal-upload-receipt.jpg").read_bytes())
    assert out.state == "partial" and "machine OCR" in out.omissions[0] and out.text


def test_failed_ocr_asset_stays_visible_in_inventory(rt, rehearsal_bundle, monkeypatch):
    from sdbackend.extract import image
    from sdbackend.extract import registry
    monkeypatch.setitem(registry.EXTRACTORS, "image/jpeg",
                        (image.IMAGE_PROFILE, lambda d: image.failed("image OCR unavailable in this runtime: test")))
    summary = _ingest(rt, rehearsal_bundle)
    assert summary["extractions"]["failed"] == 1 and summary["sources"] == 11
    assert any("extraction failed" in l and "11-05-portal-upload-receipt.jpg" in l for l in summary["limitations"])
    from sdbackend.api.ops import Operations
    ops = Operations(rt)
    ctx = ops.handle("leader.demo", dict(operation="context.resolve", preferred_view_id=None,
                                         release_id=summary["release_id"], evidence_as_of="2026-10-01T00:00:00+10:00",
                                         analysis_mode="reconstructed", analysis_as_of=None,
                                         population_reference=None))["context"]
    files = ops.handle("leader.demo", dict(operation="case.files.list", context_id=ctx["context_id"], case_id=CASE,
                                           media_types=[], cursor=None, page_size=100))
    jpg = [i for i in files["payload"]["items"] if i["asset"]["filename"].endswith(".jpg")]
    assert jpg and jpg[0]["extraction_state"] == "failed"


def test_scanned_pdf_gets_local_ocr_or_explicit_omission():
    from PIL import Image
    from sdbackend.extract.image import ocr_runtime_status
    from sdbackend.extract.pdf import extract_pdf
    buf = io.BytesIO()
    Image.open(io.BytesIO(_png_with_text("SCANNED LETTER 2026"))).save(buf, format="PDF")
    out = extract_pdf(buf.getvalue())
    if ocr_runtime_status()["available"]:
        assert out.state == "partial" and "SCANNED" in out.text.upper()
        assert any("machine OCR" in o for o in out.omissions) and out.segments[0].char_boxes
    else:
        assert out.state == "failed"


def test_nested_eml_keeps_attached_message_identity(rt, tmp_path):
    inner = email.message.EmailMessage()
    inner["From"], inner["To"], inner["Subject"] = "a@example.test", "b@example.test", "Inner message"
    inner["Date"] = "Mon, 05 Jan 2026 09:00:00 +1100"
    inner.set_content("Inner body text.")
    inner.add_attachment(b"col\n1\n", maintype="text", subtype="csv", filename="inner.csv")
    outer = email.message.EmailMessage()
    outer["From"], outer["To"], outer["Subject"] = "c@example.test", "d@example.test", "Fwd"
    outer["Date"] = "Tue, 06 Jan 2026 09:00:00 +1100"
    outer.set_content("Forwarding the message.")
    outer.add_attachment(inner)
    from sdbackend.extract.eml import extract_eml
    out = extract_eml(outer.as_bytes())
    kinds = [(c.mime_part_path, c.content_type) for c in out.children]
    assert ("0.2", "message/rfc822") in kinds and not any(p.startswith("0.2.") for p, _ in kinds)
    bundle = mini_bundle(tmp_path, [{"original_id": "ORIG-NEST", "rel": "originals/fwd.eml", "data": outer.as_bytes(),
                                     "media": "message/rfc822", "case": "C-NEST",
                                     "available": "2026-01-06T09:00:00+11:00"}])
    summary = _ingest(rt, bundle)
    media = sorted(a["media_type"] for a in rt.db.list_objects("asset"))
    assert media.count("message/rfc822") == 2 and "text/csv" in media      # inner message AND its attachment
    assert summary["attachments"] == 2


def test_same_bytes_in_two_cases_keep_both_occurrences(rt, tmp_path):
    data = b"Statement of advice, identical bytes.\n"
    bundle = mini_bundle(tmp_path, [
        {"original_id": "ORIG-A", "rel": "originals/a.txt", "data": data, "media": "text/plain", "case": "C-A",
         "available": "2026-01-06T09:00:00+11:00"},
        {"original_id": "ORIG-B", "rel": "originals/b.txt", "data": data, "media": "text/plain", "case": "C-B",
         "available": "2026-02-06T09:00:00+11:00"}])
    _ingest(rt, bundle)
    assets = rt.db.list_objects("asset")
    assert sorted(tuple(a["case_ids"]) for a in assets) == [("C-A",), ("C-B",)]
    assert len({a["bytes_sha256"] for a in assets}) == 1 and len({a["blob_reference"] for a in assets}) == 1
    assert sorted(a["scenario_available_at"] for a in assets) == ["2026-01-06T09:00:00+11:00", "2026-02-06T09:00:00+11:00"]


def test_day_precision_is_never_an_invented_instant(rt, tmp_path):
    bundle = mini_bundle(tmp_path, [{"original_id": "ORIG-DAY", "rel": "originals/d.txt", "data": b"Day-only note.\n",
                                     "media": "text/plain", "case": "C-DAY", "available": "2026-06-15",
                                     "created": "2026-06-14", "precision": "day"}])
    _ingest(rt, bundle)
    a = rt.db.list_objects("asset")[0]
    s = rt.db.list_objects("source")[0]
    assert a["scenario_document_time"] == {"precision": "day", "date": "2026-06-14"}
    assert a["scenario_available_at"] == "2026-06-16T00:00:00+10:00"          # not before the day ends (winter AEST)
    assert s["envelope"]["occurred_at"] is None and s["envelope"]["occurrence_date"] == "2026-06-14"


def test_coverage_boundaries_follow_melbourne_dst(ingested):
    rt, summary = ingested
    cov = rt.db.list_objects("coverage")[0]["data"]
    assert cov["start"] == "2025-10-01T00:00:00+10:00"        # DST begins 5 Oct 2025, so 1 Oct is still +10:00
    assert cov["end_exclusive"] == "2026-10-01T00:00:00+10:00"


def test_docx_omissions_are_declared():
    import docx
    from sdbackend.extract.textual import extract_docx
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Header text"
    p = d.add_paragraph("Col A")
    p.add_run().add_tab()
    p.add_run("Col B")
    buf = io.BytesIO()
    d.save(buf)
    out = extract_docx(buf.getvalue())
    assert out.state == "partial"
    assert any("tab/break" in o for o in out.omissions) and any("header" in o for o in out.omissions)


def test_native_rows_validated_and_procedures_retrievable(rt, rehearsal_bundle, tmp_path):
    proc = {"schema_version": "0.2.0", "kind": "procedure",
            "envelope": {"record_id": "REC-PROC-1", "logical_id": "PROC-1", "supersedes_record_id": None,
                         "recorded_at": "2025-09-01T00:00:00+10:00", "available_at": "2025-09-01T00:00:00+10:00",
                         "occurred_at": None, "occurrence_date": None,
                         "access": {"view_ids": ["complaints-leadership"], "classification": "standard",
                                    "redaction_state": "none"}, "evidence_refs": []},
            "data": {"procedure_id": "PROC-1", "version": "3", "source_id": "SRC-PROC-1", "effective_from": "2025-07-01",
                     "effective_to_exclusive": None, "process_name": "Death benefit evidence assessment"}}
    bad = case_row("SPEC-BAD", "not-a-date")
    summary, bundle = ingest_specimen(rt, tmp_path, rehearsal_bundle, [case_row("SPEC-A", "2026-01-15"), proc, bad])
    receipt = json.loads(json.dumps(__import__("sdbackend.intake.delivery", fromlist=["receive"]).receive(bundle, rt=rt).receipt))
    rows = [r for f in receipt["native_administrative_inputs"] for r in f["rows"]]
    assert any(r["state"] == "rejected" for r in rows)                      # schema-invalid row refused
    assert summary["natives"] == 2
    from sdbackend.api.repo import Repo, served_release
    from sdbackend.pipeline.tools import ToolHost, build_index
    from sdbackend.policy import pipeline_authority
    rv = served_release(rt)
    repo = Repo(rt, rv, pipeline_authority(), "2026-10-01T00:00:00+10:00", "inclusive_cutoff")
    host = ToolHost(rt, repo, run_id="RUN-t", index_version=build_index(rt, rv.release_id), case_scope=[CASE])
    found = host.execute("read_procedures", {"process_query": "death benefit evidence", "effective_date": "2026-01-01"})
    assert [i["record_id"] for i in found["items"]] == ["REC-PROC-1"]
    none = host.execute("read_procedures", {"process_query": "death benefit evidence", "effective_date": "2025-01-01"})
    assert none["items"] == [] and none["outcome"] == "no_match"


def test_candidate_facts_respect_case_filter_and_host_scope(ingested):
    rt, summary = ingested
    from sdbackend.api.repo import Repo, served_release
    from sdbackend.pipeline.tools import ToolHost, build_index
    from sdbackend.policy import pipeline_authority
    rv = served_release(rt)
    repo = Repo(rt, rv, pipeline_authority(), "2026-10-01T00:00:00+10:00", "inclusive_cutoff")
    src = repo.case_source_ids(CASE)[0]
    data = rt.db.get_object("source", src)["data"]
    from sdbackend.citations import text_ref
    ref = text_ref(data, 0, 5)
    fact = {"data": {"fact_id": "F1", "field_key": "channel", "value": {"value": "email", "state": "known",
                                                                        "evidence_refs": [ref]},
                     "asserted_by_role": "x"}}
    rt.db.put_object("stage_s02", "STAGE-T", {"facts": [fact]})
    host = ToolHost(rt, repo, run_id="RUN-t", index_version=build_index(rt, rv.release_id), case_scope=[CASE],
                    candidate_stage_ids=["STAGE-T"])
    assert len(host.execute("inspect_candidate_facts", {"case_ids": [CASE]})["items"]) == 1
    assert host.execute("inspect_candidate_facts", {"case_ids": ["C-OTHER"]})["items"] == []
    other = ToolHost(rt, repo, run_id="RUN-t", index_version=build_index(rt, rv.release_id), case_scope=["C-OTHER"],
                     candidate_stage_ids=["STAGE-T"])
    assert other.execute("inspect_candidate_facts", {})["items"] == []        # host scope, not just the argument
