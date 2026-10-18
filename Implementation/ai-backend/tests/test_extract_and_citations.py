import io

import pytest

from sdbackend.canonical import source_content_sha256
from sdbackend.citations import QuoteError, check_citation, find_quote, map_citation
from sdbackend.extract.registry import extract, validate_bytes
from sdbackend.extract.textual import tokenize_csv


def test_csv_quoted_newline_does_not_create_record():
    text = 'a,b\n"x\ny",2\n3,4\n'
    recs = tokenize_csv(text)
    assert len(recs) == 3
    s, e, v = recs[1][0]
    assert v == "x\ny" and text[s:e] == '"x\ny"'
    p, out = extract(text.encode(), "text/csv")
    assert out.state == "complete"
    second = [seg for seg in out.segments if seg.location.get("data_record_index") == 2]
    assert second and second[0].location["column_index"] == 1


def test_csv_bom_and_syntax_rejected():
    assert extract(b"\xef\xbb\xbfa,b\n1,2\n", "text/csv")[1].state == "failed"
    assert extract(b'a,b\n"x"y,2\n', "text/csv")[1].state == "failed"


def test_mime_signature_mismatch_rejected():
    ok, detected, why = validate_bytes(b"%PDF-1.4 fake", "image/png")
    assert not ok and detected == "application/pdf"
    ok, detected, _ = validate_bytes(b"a,b\n1,2\n", "text/csv")
    assert ok


def test_corrupt_pdf_fails_explicitly():
    p, out = extract(b"%PDF-1.4\nnot really", "application/pdf")
    assert out.state == "failed" and out.text == ""


def test_xlsx_formula_never_evaluated_and_cells_mapped():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws["A1"], ws["B1"], ws["B2"] = "item", 125, "=B1*2"
    buf = io.BytesIO()
    wb.save(buf)
    p, out = extract(buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert out.state == "partial"           # formula has no stored value -> omitted, not computed
    assert "250" not in out.text and "B1: 125" in out.text
    assert any("formula" in o for o in out.omissions)
    cell = [s for s in out.segments if s.location["a1_range"] == "B1"][0]
    assert out.text[cell.start:cell.end] == "125"


def test_docx_paragraph_mapping():
    import docx
    d = docx.Document()
    d.add_paragraph("Summary of call.")
    d.add_paragraph("Member asked for the missing item.")
    buf = io.BytesIO()
    d.save(buf)
    p, out = extract(buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert out.state == "complete" and out.text == "Summary of call.\nMember asked for the missing item."
    assert [s.location["paragraph_index"] for s in out.segments] == [1, 2]


def test_find_quote_zero_and_ambiguous():
    assert find_quote("abc def abc", "def") == (4, 7)
    with pytest.raises(QuoteError):
        find_quote("abc def abc", "abc")
    with pytest.raises(QuoteError):
        find_quote("abc", "xyz")
    assert find_quote("abc def abc", "abc", occurrence=2) == (8, 11)


def test_code_points_not_utf16():
    text = "😀 é item"
    s, e = find_quote(text, "item")
    assert (s, e) == (4, 8)          # emoji is one code point (two UTF-16 units)


def _source(text):
    src = {"source_id": "SRC-T", "source_family": "S03", "subtype": "email", "case_ids": ["C-1"], "afca_matter_ids": [],
           "author_role": "x", "author_reference": None, "text": text, "content_sha256": "0" * 64,
           "copy_of_source_id": None, "thread_id": None, "attachment_source_ids": [], "native_facts": [],
           "clarification_request_id": None}
    src["content_sha256"] = source_content_sha256(src)
    return src


def test_csv_citation_maps_record_column_and_header_unmapped():
    text = "request,item,status\nREQ-1,DC-11,outstanding\n"
    _, out = extract(text.encode(), "text/csv")
    src = _source(out.text)
    asset = {"asset_version_id": "AV-1", "bytes_sha256": "a" * 64}
    ext = {"extraction_id": "EXT-1", "asset_version_id": "AV-1", "bytes_sha256": "a" * 64,
           "emitted_sources": [{"source_id": src["source_id"], "content_sha256": src["content_sha256"]}]}
    s, e = find_quote(text, "outstanding")
    ref = {"source_id": src["source_id"], "content_sha256": src["content_sha256"],
           "anchor": {"kind": "text", "start_codepoint": s, "end_codepoint": e}}
    loc = map_citation(src, out.layout(), ext, asset, ref)
    assert loc["state"] == "verified" and loc["location"]["data_record_index"] == 1 and loc["location"]["column_index"] == 3
    s, e = find_quote(text, "status")
    hdr = dict(ref, anchor={"kind": "text", "start_codepoint": s, "end_codepoint": e})
    assert map_citation(src, out.layout(), ext, asset, hdr)["state"] == "unmapped"
    wrong = dict(ref, content_sha256="b" * 64)
    assert map_citation(src, out.layout(), ext, asset, wrong)["state"] == "invalid"
    other_asset = {"asset_version_id": "AV-2", "bytes_sha256": "c" * 64}
    assert map_citation(src, out.layout(), ext, other_asset, ref)["state"] == "invalid"
    cross = dict(ref, anchor={"kind": "text", "start_codepoint": 0, "end_codepoint": len(text) - 1})
    assert map_citation(src, out.layout(), ext, asset, cross)["state"] == "unmapped"


def test_check_citation_bounds_and_pointer():
    src = _source("hello")
    assert check_citation(src, {"source_id": "SRC-T", "content_sha256": src["content_sha256"],
                                "anchor": {"kind": "text", "start_codepoint": 3, "end_codepoint": 3}}) is not None
    assert check_citation(src, {"source_id": "SRC-T", "content_sha256": src["content_sha256"],
                                "anchor": {"kind": "field", "json_pointer": "/native_facts/0"}}) is not None
